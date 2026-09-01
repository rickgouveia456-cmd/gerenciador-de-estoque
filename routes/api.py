
from flask import Blueprint, render_template, request, redirect, url_for, jsonify, flash, send_file, session
from markupsafe import Markup, escape
from datetime import datetime, date, timedelta
import io, os, json, re, logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from extensions import db
from models import (agora, Almoxarifado, Item, Movimentacao, Requisicao,
    RequisicaoMestre, RequisicaoMestreItem, Usuario, Colaborador,
    Ferramenta, HistoricoFerramenta, ItemEPI, HistoricoEPI,
    AcessoExtra, PermissaoExtra)
from core import (login_required, admin_required, almoxarife_required,
    usuario_atual, flash_html, usuario_tem_acesso_almoxarifado,
    usuario_tem_acesso_item, PERMISSOES_DISPONIVEIS,
    _check_rate_limit, _register_attempt, _clear_attempts, _check_api_rate,
    is_admin_ou_ggo)

logger = logging.getLogger(__name__)
api_bp = Blueprint('api_bp', __name__)

@api_bp.route('/api/alertas')
@login_required
def api_alertas():
    if _check_api_rate(request.remote_addr or '0.0.0.0'):
        return jsonify({'error': 'Too many requests'}), 429
    u = usuario_atual()
    if is_admin_ou_ggo(u):
        if u.perfil == 'ggo':
            from core import almoxarifados_do_ggo
            ids_ggo = {a.id for a in almoxarifados_do_ggo(u)}
            itens = Item.query.filter(
                Item.quantidade <= Item.estoque_minimo,
                Item.almoxarifado_id.in_(ids_ggo),
                Item.ativo == True
            ).all() if ids_ggo else []
        else:
            itens = Item.query.filter(Item.quantidade <= Item.estoque_minimo, Item.ativo == True).all()
    else:
        ids = u.almoxarifados_permitidos()
        itens = Item.query.filter(
            Item.quantidade <= Item.estoque_minimo,
            Item.almoxarifado_id.in_(ids),
            Item.ativo == True
        ).all() if ids else []
    return jsonify([{
        'id': i.id, 'nome': i.nome, 'codigo': i.codigo,
        'quantidade': i.quantidade, 'estoque_minimo': i.estoque_minimo,
        'unidade': i.unidade, 'status': i.status,
        'almoxarifado': i.almoxarifado.nome
    } for i in itens])

@api_bp.route('/api/almoxarife/notificacoes')
@login_required
def api_almoxarife_notificacoes():
    """Retorna requisições pendentes para o almoxarife logado — usado para popup de alerta."""
    u = usuario_atual()
    if not is_admin_ou_ggo(u) and u.perfil != 'almoxarife':
        return jsonify([])
    # Busca requisições pendentes do almoxarifado do almoxarife
    if u.perfil == 'almoxarife' and u.almoxarifado_id:
        reqs = RequisicaoMestre.query.filter_by(
            almoxarifado_id=u.almoxarifado_id,
            status='pendente'
        ).order_by(RequisicaoMestre.data_criacao.desc()).limit(5).all()
    else:
        reqs = RequisicaoMestre.query.filter_by(
            status='pendente'
        ).order_by(RequisicaoMestre.data_criacao.desc()).limit(5).all()
    return jsonify([{
        'id': r.id,
        'protocolo': r.protocolo or f'#{r.id}',
        'mestre': r.mestre.nome,
        'colaborador': r.colaborador,
        'almoxarifado': r.almoxarifado.nome,
        'itens': len(r.itens),
        'data': r.data_criacao.strftime('%d/%m/%Y %H:%M')
    } for r in reqs])

@api_bp.route('/api/colaboradores')
@login_required
def api_colaboradores():
    if _check_api_rate(request.remote_addr or '0.0.0.0'):
        return jsonify([]), 429
    """Autocomplete — busca colaboradores por nome, filtrado por cidade do usuário."""
    q = request.args.get('q', '').strip()
    u = usuario_atual()
    nomes = []
    like = f"%{q}%"
    is_pg = 'postgresql' in str(db.engine.url)
    ilike = 'ILIKE' if is_pg else 'LIKE'

    from sqlalchemy import text

    # Determinar cidade e escopo do usuário logado (via almoxarifado vinculado)
    escopo_usuario = None
    cidade_usuario = None
    if u and u.almoxarifado_id:
        alm = db.session.get(Almoxarifado, u.almoxarifado_id)
        if alm:
            cidade_usuario = (alm.cidade or '').strip() or None
            nome_lower = alm.nome.lower()
            for esc in ['estrutura', 'acabamento', 'infraestrutura', 'forma', 'acampamento']:
                if esc in nome_lower:
                    escopo_usuario = esc
                    break

    # Filtro de cidade — admin vê todos, demais só da sua cidade
    cidade_filtro = cidade_usuario if (u and u.perfil != 'admin') else None

    def _query_colabs(extra_where='', extra_params=None):
        params = {"q": like}
        cidade_clause = ''
        if cidade_filtro:
            cidade_clause = f" AND (cidade {ilike} :cidade OR cidade IS NULL OR cidade = '')" if not extra_where else \
                            f" AND (cidade {ilike} :cidade OR cidade IS NULL OR cidade = '')"
            params["cidade"] = cidade_filtro
        if extra_params:
            params.update(extra_params)
        sql = f"SELECT nome, funcao, escopo, tipo FROM colaborador WHERE ativo = TRUE AND nome {ilike} :q{cidade_clause}{extra_where} ORDER BY nome LIMIT 15"
        return db.session.execute(text(sql), params).fetchall()

    # 1. Colaboradores do mesmo escopo primeiro
    if escopo_usuario:
        rows = _query_colabs(
            extra_where=f" AND escopo {ilike} :esc",
            extra_params={"esc": f"%{escopo_usuario}%"}
        )
        for r in rows:
            nomes.append({'nome': r[0], 'funcao': r[1] or '', 'escopo': r[2] or '', 'tipo': r[3] or 'peao'})

    # 2. Demais colaboradores da mesma cidade
    nomes_ja = {n['nome'] for n in nomes}
    rows = _query_colabs()
    for r in rows:
        if r[0] not in nomes_ja:
            nomes.append({'nome': r[0], 'funcao': r[1] or '', 'escopo': r[2] or '', 'tipo': r[3] or 'peao'})

    # 3. Histórico de requisições dos almoxarifados da cidade (fallback)
    nomes_ja = {n['nome'] for n in nomes}
    if cidade_filtro:
        # Só requisições de almoxarifados da mesma cidade
        ids_cidade = [a.id for a in Almoxarifado.query.filter(
            Almoxarifado.cidade.ilike(cidade_filtro)
        ).all()]
        if ids_cidade:
            placeholders = ','.join(str(i) for i in ids_cidade)
            rows = db.session.execute(
                text(f"SELECT DISTINCT colaborador FROM requisicao_mestre WHERE colaborador {ilike} :q AND almoxarifado_id IN ({placeholders}) ORDER BY colaborador LIMIT 5"),
                {"q": like}
            ).fetchall()
            for r in rows:
                if r[0] not in nomes_ja:
                    nomes.append({'nome': r[0], 'funcao': '', 'escopo': '', 'tipo': ''})
    else:
        rows = db.session.execute(
            text(f"SELECT DISTINCT colaborador FROM requisicao_mestre WHERE colaborador {ilike} :q ORDER BY colaborador LIMIT 5"),
            {"q": like}
        ).fetchall()
        for r in rows:
            if r[0] not in nomes_ja:
                nomes.append({'nome': r[0], 'funcao': '', 'escopo': '', 'tipo': ''})

    return jsonify(nomes[:12])


@api_bp.route('/api/epis/buscar')
@login_required
def api_epis_buscar():
    """Autocomplete de EPIs — busca em ItemEPI, Item(categoria=epi) e CatalogoInsumo."""
    from models import CatalogoInsumo, ItemFichaEPI
    q = request.args.get('q', '').strip()
    u = usuario_atual()
    if not q or len(q) < 1:
        return jsonify([])

    is_pg = 'postgresql' in str(db.engine.url)
    ilike = 'ILIKE' if is_pg else 'LIKE'
    like = f'%{q}%'

    resultados = set()

    # 1. ItemEPI dos almoxarifados do usuário
    if u.perfil == 'admin':
        epis = ItemEPI.query.filter(
            ItemEPI.nome.ilike(like), ItemEPI.ativo == True
        ).order_by(ItemEPI.nome).limit(8).all()
    else:
        ids = u.almoxarifados_permitidos()
        epis = ItemEPI.query.filter(
            ItemEPI.nome.ilike(like),
            ItemEPI.ativo == True,
            ItemEPI.almoxarifado_id.in_(ids)
        ).order_by(ItemEPI.nome).limit(8).all() if ids else []

    for e in epis:
        resultados.add(e.nome)

    # 2. Item com categoria=epi
    if u.perfil == 'admin':
        itens = Item.query.filter(
            Item.nome.ilike(like), Item.categoria == 'epi', Item.ativo == True
        ).order_by(Item.nome).limit(6).all()
    else:
        ids = u.almoxarifados_permitidos()
        itens = Item.query.filter(
            Item.nome.ilike(like), Item.categoria == 'epi', Item.ativo == True,
            Item.almoxarifado_id.in_(ids)
        ).order_by(Item.nome).limit(6).all() if ids else []

    for it in itens:
        resultados.add(it.nome)

    # 3. Catálogo de insumos com categoria=epi
    cat_epis = CatalogoInsumo.query.filter(
        CatalogoInsumo.nome.ilike(like),
        CatalogoInsumo.categoria == 'epi',
        CatalogoInsumo.ativo == True
    ).order_by(CatalogoInsumo.nome).limit(6).all()
    for c in cat_epis:
        resultados.add(c.nome)

    # 4. Histórico de descrições já usadas em fichas
    try:
        rows = db.session.execute(
            db.text(f"SELECT DISTINCT descricao FROM item_ficha_epi WHERE descricao {ilike} :q ORDER BY descricao LIMIT 6"),
            {'q': like}
        ).fetchall()
        for r in rows:
            resultados.add(r[0])
    except Exception:
        pass

    lista = sorted(resultados)[:15]
    return jsonify([{'nome': n} for n in lista])

# ── FROTA DE FERRAMENTAS ─────────────────────────────────────────────────────


@api_bp.route('/api/mestre/notificacoes')
@login_required
def api_mestre_notificacoes():
    """Retorna notificações baseadas no status das requisições do mestre."""
    u = usuario_atual()
    if u.perfil not in ('mestre', 'tecnico_seguranca'):
        return jsonify([])
    # Busca requisições não-pendentes e não-canceladas do mestre
    # A notificação é baseada no status atual — o frontend controla o que já foi visto via localStorage
    reqs = RequisicaoMestre.query.filter_by(mestre_id=u.id).filter(
        RequisicaoMestre.status.in_(['aprovada', 'parcial', 'recusada', 'entregue'])
    ).order_by(RequisicaoMestre.data_criacao.desc()).limit(10).all()

    result = []
    for r in reqs:
        if r.status == 'aprovada':
            msg = f'✅ Requisição #{r.id} aprovada! Envie o colaborador buscar.'
            tipo = 'success'
        elif r.status == 'parcial':
            msg = f'⚠️ Requisição #{r.id} aprovada parcialmente. Verifique os itens.'
            tipo = 'warning'
        elif r.status == 'recusada':
            msg = f'❌ Requisição #{r.id} foi recusada pelo almoxarifado.'
            tipo = 'danger'
        elif r.status == 'entregue':
            msg = f'📦 Requisição #{r.id} entregue ao colaborador {r.colaborador}.'
            tipo = 'info'
        else:
            continue
        result.append({'id': r.id, 'msg': msg, 'tipo': tipo, 'status': r.status})
    return jsonify(result)

@api_bp.route('/api/mestre/notificacoes/marcar-lidas', methods=['POST'])
@login_required
def marcar_notificacoes_lidas():
    """Endpoint mantido para compatibilidade — controle feito no frontend."""
    return jsonify({'ok': True})

# ── ROTA ESPECIAL PARA REATIVAR TODOS OS ITENS ──────────────────────────────
